from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\HP\Documents\Project end\SmartLifeExpo")
OUTPUT = ROOT / "SmartLife_Sequence_Diagrams.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(20 if level == 1 else 16)
    run.font.color.rgb = RGBColor(40, 65, 36)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(12)
    return paragraph


def add_sequence_table(document, steps):
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0].cells
    for cell, label in zip(header, ["ลำดับ", "ผู้ส่ง", "ผู้รับ", "ข้อความ / การทำงาน"]):
        set_cell_shading(cell, "4F794C")
        set_cell_text(cell, label, bold=True, color="FFFFFF", size=11)
    for index, (sender, receiver, action) in enumerate(steps, start=1):
        cells = table.add_row().cells
        values = [str(index), sender, receiver, action]
        for column, (cell, value) in enumerate(zip(cells, values)):
            if index % 2 == 0:
                set_cell_shading(cell, "F1F5EE")
            set_cell_text(cell, value, size=10)
            if column == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()


def add_participants(document, participants):
    paragraph = document.add_paragraph()
    label = paragraph.add_run("องค์ประกอบหลัก: ")
    label.bold = True
    label.font.name = "TH Sarabun New"
    label._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    label.font.size = Pt(12)
    value = paragraph.add_run("  |  ".join(participants))
    value.font.name = "TH Sarabun New"
    value._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    value.font.size = Pt(12)


def add_diagram_section(document, title, overview, participants, steps):
    add_heading(document, title, 1)
    add_body(document, overview)
    add_participants(document, participants)
    add_sequence_table(document, steps)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SmartLife - Sequence Diagrams")
    run.bold = True
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(40, 65, 36)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("เอกสารลำดับการทำงานของระบบ SmartLife")
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(102, 120, 99)
    document.add_paragraph()

    add_body(document, "เอกสารนี้สรุปการสื่อสารระหว่างผู้ใช้ แอป SmartLife และบริการภายนอก เพื่อใช้ประกอบการพัฒนา การทดสอบ และการนำเสนอโครงงานเทคโนโลยีดิจิทัล")

    add_diagram_section(
        document,
        "1. Login, Role และ Session",
        "ผู้ใช้หรือผู้ดูแลระบบเข้าสู่ระบบด้วยอีเมล/รหัสผ่าน หรือ Google Login ระบบตรวจสอบตัวตนและส่งผู้ใช้ไปยังหน้าตามบทบาท พร้อมเก็บสถานะการเข้าใช้งานอย่างปลอดภัย",
        ["User / Admin", "SmartLife Expo App", "Firebase Authentication", "Cloud Firestore", "Secure Session Storage"],
        [
            ("User / Admin", "SmartLife Expo App", "เปิดแอปและตรวจสอบ session เดิม"),
            ("SmartLife Expo App", "Firebase Authentication", "อ่าน currentUser"),
            ("Firebase Authentication", "SmartLife Expo App", "ส่ง uid หรือ null"),
            ("SmartLife Expo App", "Cloud Firestore", "อ่าน users/{uid} เพื่อรับ profile และ role"),
            ("Cloud Firestore", "SmartLife Expo App", "ตอบข้อมูลผู้ใช้และ role"),
            ("SmartLife Expo App", "Secure Session Storage", "บันทึกสถานะ session"),
            ("SmartLife Expo App", "User / Admin", "เปิด User Dashboard หรือ Admin Dashboard ตาม role"),
        ],
    )

    add_diagram_section(
        document,
        "2. Unified Dashboard และ Dynamic Prioritization",
        "หน้าแรกดึงตารางเรียน งาน โน้ต การเงิน และคำแนะนำ AI ของผู้ใช้คนปัจจุบัน แล้วจัดลำดับสิ่งที่ต้องทำด่วนขึ้นก่อน",
        ["User", "SmartLife App", "Cloud Firestore", "AI Recommendation Function"],
        [
            ("User", "SmartLife App", "เปิด Unified Dashboard"),
            ("SmartLife App", "Cloud Firestore", "อ่าน schedules, notes, transactions และ tasks ของ uid"),
            ("Cloud Firestore", "SmartLife App", "ตอบข้อมูลล่าสุด"),
            ("SmartLife App", "AI Recommendation Function", "ส่งบริบทของวันนี้เพื่อจัดลำดับ"),
            ("AI Recommendation Function", "SmartLife App", "ตอบรายการสำคัญและคำแนะนำ"),
            ("SmartLife App", "User", "แสดงการ์ดด่วน ข้อมูลสรุป และ AI Assistant"),
        ],
    )

    add_diagram_section(
        document,
        "3. Smart Schedule Importer และ OCR ตารางเรียน",
        "ผู้ใช้อัปโหลดรูปตารางเรียน ระบบเก็บภาพอย่างปลอดภัย ประมวลผล OCR และ AI จากนั้นเปิดหน้าตรวจแก้ก่อนบันทึกตารางแบบเกิดซ้ำรายสัปดาห์",
        ["User", "SmartLife App", "Cloud Storage", "Cloud Function", "Vision OCR / Gemini", "Cloud Firestore"],
        [
            ("User", "SmartLife App", "เลือกกล้องหรือคลังภาพ"),
            ("SmartLife App", "Cloud Storage", "อัปโหลดภาพตาม uid/scanId"),
            ("SmartLife App", "Cloud Function", "เรียก analyzeScan พร้อม storage path"),
            ("Cloud Function", "Vision OCR / Gemini", "อ่านข้อความและจำแนกเป็นตารางเรียน"),
            ("Vision OCR / Gemini", "Cloud Function", "ส่งวิชา วัน เวลา ห้อง และข้อมูลสอบ"),
            ("Cloud Function", "SmartLife App", "ส่ง structured schedule data"),
            ("User", "SmartLife App", "แก้ไขข้อมูลและยืนยัน"),
            ("SmartLife App", "Cloud Firestore", "บันทึก schedules พร้อม seriesId และ scan log"),
        ],
    )

    add_diagram_section(
        document,
        "4. Receipt / Slip OCR และ Smart Expense Categorization",
        "ผู้ใช้สแกนใบเสร็จหรือสลิป ระบบวิเคราะห์ชื่อร้าน รายการสินค้า ยอดรวม วันเวลา และจัดหมวดหมู่รายจ่ายก่อนให้ผู้ใช้ตรวจสอบและบันทึก",
        ["User", "SmartLife App", "Cloud Storage", "Cloud Function", "Vision OCR / Gemini", "Cloud Firestore"],
        [
            ("User", "SmartLife App", "ถ่ายหรือเลือกรูปใบเสร็จ/สลิป"),
            ("SmartLife App", "Cloud Storage", "อัปโหลดภาพสแกน"),
            ("SmartLife App", "Cloud Function", "เรียก analyzeScan"),
            ("Cloud Function", "Vision OCR / Gemini", "อ่าน merchant, items, total, date และ time"),
            ("Cloud Function", "Cloud Function", "จัดหมวด Food, Groceries, Utilities, Transport, Entertainment, Shopping หรือ Others"),
            ("Cloud Function", "SmartLife App", "ส่งผล OCR พร้อมรายการสินค้า"),
            ("User", "SmartLife App", "ตรวจแก้รายการและยอดรวม"),
            ("SmartLife App", "Cloud Firestore", "บันทึก transactions และ scanLogs"),
        ],
    )

    add_diagram_section(
        document,
        "5. Google Calendar Two-Way Sync",
        "ผู้ใช้เชื่อมบัญชี Google ด้วย OAuth เพื่อดึงกิจกรรมจาก Google Calendar และส่งตารางเรียนที่บันทึกแล้วกลับไปยัง Google Calendar",
        ["User", "SmartLife App", "Google OAuth", "Google Calendar API", "Cloud Firestore"],
        [
            ("User", "SmartLife App", "กดเชื่อม Google Calendar"),
            ("SmartLife App", "Google OAuth", "ขอสิทธิ์ profile และ calendar.events"),
            ("Google OAuth", "SmartLife App", "ตอบ OAuth token และข้อมูลบัญชี"),
            ("SmartLife App", "Cloud Firestore", "เก็บสถานะการเชื่อมต่อ calendarConnections/{uid}"),
            ("SmartLife App", "Google Calendar API", "ดึงกิจกรรมในช่วงวันที่กำหนด"),
            ("Google Calendar API", "SmartLife App", "ตอบรายการกิจกรรม"),
            ("SmartLife App", "Google Calendar API", "เพิ่ม/อัปเดตกิจกรรมตารางเรียนแบบ recurring"),
        ],
    )

    add_diagram_section(
        document,
        "6. AI Assistant, Adaptive Scheduling และ Burnout Coach",
        "AI Assistant รับคำถามด้วยข้อความหรือเสียง วิเคราะห์ข้อมูลส่วนตัวของผู้ใช้ และตอบคำแนะนำเรื่องตารางเรียน การเงิน งาน หรือความเสี่ยง burnout",
        ["User", "SmartLife App", "Cloud Firestore", "Cloud Function", "Gemini API"],
        [
            ("User", "SmartLife App", "ส่งคำถามหรือคำสั่งเสียง"),
            ("SmartLife App", "Cloud Firestore", "อ่านบริบทตาราง โน้ต การเงิน และงานของ uid"),
            ("SmartLife App", "Cloud Function", "ส่ง prompt พร้อมบริบทที่ผ่านการคัดเลือก"),
            ("Cloud Function", "Gemini API", "ขอคำตอบหรือคำแนะนำแบบมีโครงสร้าง"),
            ("Gemini API", "Cloud Function", "ส่งคำตอบและ priority recommendation"),
            ("Cloud Function", "Cloud Firestore", "บันทึก aiRecommendations และ audit log"),
            ("SmartLife App", "User", "แสดงคำตอบ คำแนะนำตาราง และ burnout coaching"),
        ],
    )

    add_diagram_section(
        document,
        "7. Admin Dashboard และการดูแลระบบ",
        "ผู้ดูแลระบบดูข้อมูลภาพรวมจาก Firebase เช่น จำนวนผู้ใช้ ตารางเรียน โน้ต รายการการเงิน การสแกน และสถานะบริการ โดยไม่เปิดเผยข้อมูลส่วนตัวเกินสิทธิ์",
        ["Admin", "Admin App", "Firebase Authentication", "Cloud Function", "Cloud Firestore"],
        [
            ("Admin", "Admin App", "เข้าสู่ระบบ"),
            ("Admin App", "Firebase Authentication", "ตรวจสอบตัวตนและ admin claim"),
            ("Admin App", "Cloud Function", "ขอข้อมูลสรุปที่ต้องใช้สิทธิ์ผู้ดูแล"),
            ("Cloud Function", "Cloud Firestore", "นับและอ่านข้อมูลที่อนุญาต"),
            ("Cloud Firestore", "Cloud Function", "ตอบสถิติ health logs feedback และ scan logs"),
            ("Cloud Function", "Admin App", "ตอบ dashboard metrics และรายการดูแลระบบ"),
            ("Admin App", "Admin", "แสดงผล dashboard, feedback, OCR log และ system health"),
        ],
    )

    add_heading(document, "8. โครงสร้างข้อมูลหลัก", 1)
    add_body(document, "ตารางต่อไปนี้สรุป collection หลักของ Cloud Firestore ที่ใช้เชื่อมทุกหน้าของ SmartLife")
    data_table = document.add_table(rows=1, cols=3)
    data_table.style = "Table Grid"
    for cell, label in zip(data_table.rows[0].cells, ["Collection", "หน้าที่", "ตัวอย่างข้อมูลสำคัญ"]):
        set_cell_shading(cell, "4F794C")
        set_cell_text(cell, label, bold=True, color="FFFFFF", size=11)
    records = [
        ("users/{uid}", "โปรไฟล์และบทบาท", "displayName, email, role, institutionType, calendarConnected"),
        ("schedules", "ตารางเรียนและกิจกรรม", "uid, courseCode, courseName, startAt, endAt, seriesId"),
        ("transactions", "รายรับรายจ่าย", "uid, merchantName, items, totalAmount, category, occurredAt"),
        ("notes", "โน้ตของผู้ใช้", "uid, title, content, category, linkedScheduleId"),
        ("scanLogs", "ประวัติการสแกน OCR", "uid, scanType, imagePath, rawText, confidence, result"),
        ("aiRecommendations", "คำแนะนำของ AI", "uid, prompt, response, priority, sourceReferences"),
        ("calendarConnections", "สถานะเชื่อม Google Calendar", "uid, provider, email, connectedAt, status"),
        ("feedback", "ข้อเสนอแนะจากผู้ใช้", "uid, category, message, status, createdAt"),
        ("announcements", "ประกาศจากผู้ดูแล", "title, content, audience, publishedAt"),
    ]
    for index, record in enumerate(records):
        cells = data_table.add_row().cells
        for cell, value in zip(cells, record):
            if index % 2 == 1:
                set_cell_shading(cell, "F1F5EE")
            set_cell_text(cell, value, size=10)

    document.add_paragraph()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("SmartLife Senior Project - Sequence Diagrams")
    footer_run.font.name = "TH Sarabun New"
    footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    footer_run.font.size = Pt(10)
    footer_run.font.color.rgb = RGBColor(102, 120, 99)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
