from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SMARTLIFE_SEQUENCE_DIAGRAMS.md"
OUTPUT = ROOT / "SMARTLIFE_SEQUENCE_DIAGRAMS_1_TO_DATA_STRUCTURE.docx"


def markdown_heading(line: str) -> tuple[int, str] | None:
    match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
    if not match:
        return None
    return len(match.group(1)), match.group(2).strip()


def choose_requested_content(lines: list[str]) -> list[str]:
    headings = [(index, *value) for index, line in enumerate(lines) if (value := markdown_heading(line))]
    start = next(
        (index for index, _level, title in headings if re.match(r"1[.)]\s", title)),
        0,
    )
    target = next(
        ((index, level) for index, level, title in headings if "โครงสร้างข้อมูลหลัก" in title),
        None,
    )
    if not target:
        return lines[start:]

    target_index, target_level = target
    end = len(lines)
    for index, level, _title in headings:
        if index > target_index and level <= target_level:
            end = index
            break
    return lines[start:end]


def set_font(style, name: str, size: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    set_font(document.styles["Normal"], "TH Sarabun New", 13)
    for name, size in (("Title", 24), ("Heading 1", 19), ("Heading 2", 16), ("Heading 3", 14)):
        set_font(document.styles[name], "TH Sarabun New", size, True)

    header = section.header.paragraphs[0]
    header.text = "SmartLife | Sequence Diagrams"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(105, 130, 98)

    footer = section.footer.paragraphs[0]
    footer.text = "เอกสารประกอบโครงงาน SmartLife"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(125, 125, 125)


def render_mermaid(source: str, stem: str, work_dir: Path) -> Path | None:
    input_path = work_dir / f"{stem}.mmd"
    output_path = work_dir / f"{stem}.png"
    puppeteer_config = work_dir / "puppeteer.json"
    input_path.write_text(source, encoding="utf-8")
    puppeteer_config.write_text('{"args":["--no-sandbox","--disable-setuid-sandbox"]}', encoding="utf-8")

    local_cli = ROOT / "node_modules" / ".bin" / "mmdc.cmd"
    if local_cli.exists():
        command = [str(local_cli)]
    elif shutil.which("mmdc.cmd"):
        command = ["mmdc.cmd"]
    else:
        command = ["npx.cmd", "--yes", "@mermaid-js/mermaid-cli@11.4.2"]

    command += [
        "-i", str(input_path),
        "-o", str(output_path),
        "-b", "transparent",
        "-w", "1700",
        "-p", str(puppeteer_config),
    ]
    try:
        subprocess.run(
            command,
            cwd=ROOT,
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=120,
        )
    except (subprocess.SubprocessError, OSError):
        return None
    return output_path if output_path.exists() else None


def add_code_fallback(document: Document, mermaid_source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    run = paragraph.add_run(mermaid_source.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def build_document() -> tuple[int, int]:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source not found: {SOURCE}")

    content = choose_requested_content(SOURCE.read_text(encoding="utf-8").splitlines())
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("SmartLife - Sequence Diagrams")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("ลำดับการทำงานของระบบ ตั้งแต่ข้อ 1 ถึงโครงสร้างข้อมูลหลัก")
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.color.rgb = RGBColor(105, 130, 98)

    diagrams = 0
    fallback_diagrams = 0
    heading_count = 0
    lines = iter(enumerate(content))
    with tempfile.TemporaryDirectory(prefix="smartlife_mermaid_") as directory:
        work_dir = Path(directory)
        for _index, line in lines:
            heading = markdown_heading(line)
            if heading:
                level, text = heading
                if heading_count and level == 1:
                    document.add_page_break()
                document.add_heading(text, level=min(level, 3))
                heading_count += 1
                continue

            if line.strip().startswith("```mermaid"):
                mermaid_lines: list[str] = []
                for _diagram_index, diagram_line in lines:
                    if diagram_line.strip() == "```":
                        break
                    mermaid_lines.append(diagram_line)
                diagram_source = "\n".join(mermaid_lines).strip()
                if diagram_source:
                    image_path = render_mermaid(diagram_source, f"sequence_{diagrams + fallback_diagrams + 1}", work_dir)
                    if image_path:
                        picture = document.add_picture(str(image_path), width=Inches(10.75))
                        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        diagrams += 1
                    else:
                        add_code_fallback(document, diagram_source)
                        fallback_diagrams += 1
                continue

            text = line.strip()
            if not text:
                continue
            if text.startswith("-") or text.startswith("*"):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(text[1:].strip())
            elif re.match(r"^\d+[.)]\s", text):
                paragraph = document.add_paragraph(style="List Number")
                paragraph.add_run(re.sub(r"^\d+[.)]\s*", "", text))
            else:
                paragraph = document.add_paragraph(text)
                paragraph.paragraph_format.space_after = Pt(5)

    document.save(OUTPUT)
    return diagrams, fallback_diagrams


if __name__ == "__main__":
    rendered, fallback = build_document()
    print(f"Created: {OUTPUT}")
    print(f"Mermaid images: {rendered}; source fallbacks: {fallback}")
